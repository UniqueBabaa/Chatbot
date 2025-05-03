import streamlit as st
import PyPDF2
import docx  # for reading Word files
from langchain.embeddings import OpenAIEmbeddings
from langchain.vectorstores import FAISS
from langchain.chat_models import ChatOpenAI
from langchain.chains.qa_with_sources import load_qa_with_sources_chain
from langchain.docstore.document import Document

## you have to provide the OpenAI API key below
OPENAI_API_KEY = "##provide your open ai api key here"  

## extract text from uploaded PDFs and DOCX files with metadata
def load_documents(files):
    documents = []
    for file in files:
        file_name = file.name
        if file_name.lower().endswith(".pdf"):
            reader = PyPDF2.PdfReader(file)
            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    doc = Document(
                        page_content=text,
                        metadata={"source": f"{file_name} - Page {i + 1}"}
                    )
                    documents.append(doc)
        elif file_name.lower().endswith(".docx"):
            docx_file = docx.Document(file)
            full_text = "\n".join([para.text for para in docx_file.paragraphs if para.text.strip()])
            if full_text:
                doc = Document(
                    page_content=full_text,
                    metadata={"source": f"{file_name}"}
                )
                documents.append(doc)
    return documents

def create_vectorstore(documents):
    embeddings = OpenAIEmbeddings(openai_api_key=OPENAI_API_KEY)
    texts = [doc.page_content for doc in documents]
    metadatas = [doc.metadata for doc in documents]
    vectorstore = FAISS.from_texts(texts, embeddings, metadatas=metadatas)
    return vectorstore

def build_qa_chain(vectorstore):
    retriever = vectorstore.as_retriever()
    llm = ChatOpenAI(model="gpt-4", openai_api_key=OPENAI_API_KEY)

    def custom_qa_chain(question):
        docs = retriever.get_relevant_documents(question)
        if not docs:
            return "Sorry, I don't have relevant information for this question based on the provided document(s)."
        qa_chain = load_qa_with_sources_chain(llm, chain_type="stuff")
        result = qa_chain.run(input_documents=docs, question=question)
        return result

    return custom_qa_chain

# Streamlit UI
st.set_page_config(page_title="PDF & Word Chatbot", layout="wide")
st.title("📄 Chatbot on Customer Support Documentation")

uploaded_files = st.file_uploader("Upload PDF or Word (.docx) files", type=["pdf", "docx"], accept_multiple_files=True)

if uploaded_files:
    with st.spinner("Reading and indexing documents..."):
        documents = load_documents(uploaded_files)
        vectorstore = create_vectorstore(documents)
        qa_chain = build_qa_chain(vectorstore)
    st.success("Documents loaded successfully!")

    question = st.text_input("Ask a question about the uploaded documents:")

    if question:
        with st.spinner("Thinking..."):
            response = qa_chain(question)
            if "SOURCES:" in response:
                response = response.split("SOURCES:")[0].strip()

        st.markdown("### Answer")
        st.write(response)

else:
    st.info("Please upload at least one PDF or Word document to get started.")
print ("App is running...")